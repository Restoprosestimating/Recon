"""Build the Restopros Recon Tracker workbook from parsed Internal TAM data."""
from __future__ import annotations

from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from parser import ParseResult

NAVY, TEAL = "1B365D", "0D7377"
LT_NAVY, LT_TEAL, LT_GOLD, YELLOW = "E8EEF4", "E6F3F3", "F8F1DC", "FFF2CC"
GREEN_F, RED_F, WHITE = "C6EFCE", "F8CBAD", "FFFFFF"

thin = Border(
    left=Side(style="thin", color="B0B8C1"),
    right=Side(style="thin", color="B0B8C1"),
    top=Side(style="thin", color="B0B8C1"),
    bottom=Side(style="thin", color="B0B8C1"),
)
cur = '$#,##0.00;($#,##0.00);"-"'
pct = '0.0%;(0.0%);"-"'
hrs = "0.00"


def _header(ws, row, start, end, fill):
    for c in range(start, end + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = thin


def build_workbook(parsed: ParseResult, settings: dict, path: str) -> str:
    rate_sup = float(settings.get("rate_sup", 30))
    rate_wkr = float(settings.get("rate_wkr", 22))
    burden = float(settings.get("burden", 0.18))
    royalty = float(settings.get("royalty", 0.08))
    contig = float(settings.get("contingency", 0.05))
    tax = float(settings.get("tax", 0.0825))
    negotiated = float(settings.get("negotiated") or parsed.rcv or parsed.line_item_total or 0)
    rcv = parsed.rcv or negotiated
    job = parsed.property_line or parsed.insured or "Restopros job"
    est = parsed.estimate_id or "—"

    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"

    fill_navy = PatternFill("solid", fgColor=NAVY)
    fill_teal = PatternFill("solid", fgColor=TEAL)
    fill_lt_navy = PatternFill("solid", fgColor=LT_NAVY)
    fill_lt_teal = PatternFill("solid", fgColor=LT_TEAL)
    fill_lt_gold = PatternFill("solid", fgColor=LT_GOLD)
    fill_yellow = PatternFill("solid", fgColor=YELLOW)
    fill_green = PatternFill("solid", fgColor=GREEN_F)
    fill_red = PatternFill("solid", fgColor=RED_F)
    font_title = Font(name="Calibri", size=18, bold=True, color=WHITE)
    font_body = Font(name="Calibri", size=10)
    font_small = Font(name="Calibri", size=8, italic=True, color="555555")
    font_h2 = Font(name="Calibri", size=13, bold=True, color=NAVY)
    font_white = Font(name="Calibri", size=10, bold=True, color=WHITE)
    font_input = Font(name="Calibri", size=16, bold=True, color="0000FF")

    # ----- Assumptions -----
    a = wb.create_sheet("Assumptions")
    a.merge_cells("A1:C1")
    a["A1"] = "RESTOPROS RECON TRACKER  •  ASSUMPTIONS"
    a["A1"].font = font_title
    a["A1"].fill = fill_navy
    a.row_dimensions[1].height = 26
    a["A3"] = "Job / property"
    a["B3"] = job
    a["A4"] = "Estimate #"
    a["B4"] = est
    a["A5"] = "Price list"
    a["B5"] = parsed.price_list
    labels = [
        (7, "Supervisor gross $/hr", rate_sup, cur, "CLN-S / supervisory hours"),
        (8, "Labor worker gross $/hr", rate_wkr, cur, "All other Internal TAM labor codes"),
        (9, "Payroll burden rate", burden, pct, "Employer burden on gross wages"),
        (10, "Royalty % of negotiated price", royalty, pct, "Restopros royalty / override"),
        (11, "Contingency % (matl + equip)", contig, pct, "Optional cushion"),
        (12, "Sales tax rate", tax, pct, "Used only if you treat tax as cash cost"),
    ]
    a["A6"] = "Assumption"
    a["B6"] = "Value"
    a["C6"] = "Notes"
    _header(a, 6, 1, 3, fill_navy)
    for r, lab, val, fmt, note in labels:
        a.cell(r, 1, lab).font = font_body
        a.cell(r, 1).fill = fill_lt_navy
        a.cell(r, 2, val).font = Font(name="Calibri", size=10, color="0000FF")
        a.cell(r, 2).fill = fill_yellow
        a.cell(r, 2).number_format = fmt
        a.cell(r, 3, note).font = font_small
        for c in range(1, 4):
            a.cell(r, c).border = thin
    a.column_dimensions["A"].width = 36
    a.column_dimensions["B"].width = 16
    a.column_dimensions["C"].width = 48

    wb.defined_names.add(DefinedName(name="Rate_Sup", attr_text="Assumptions!$B$7"))
    wb.defined_names.add(DefinedName(name="Rate_Wkr", attr_text="Assumptions!$B$8"))
    wb.defined_names.add(DefinedName(name="Burden_Pct", attr_text="Assumptions!$B$9"))
    wb.defined_names.add(DefinedName(name="Royalty_Pct", attr_text="Assumptions!$B$10"))
    wb.defined_names.add(DefinedName(name="Contingency_Pct", attr_text="Assumptions!$B$11"))
    wb.defined_names.add(DefinedName(name="Tax_Pct", attr_text="Assumptions!$B$12"))
    wb.defined_names.add(DefinedName(name="Billable_Amt", attr_text="Summary!$B$4"))

    # ----- Labor -----
    lab = wb.create_sheet("Labor Budget")
    lab.merge_cells("A1:L1")
    lab["A1"] = "LABOR BUDGET  •  Internal TAM hours recosted at Restopros rates"
    lab["A1"].font = font_title
    lab["A1"].fill = fill_navy
    lab.row_dimensions[1].height = 24
    headers = ["Code", "Trade", "Role", "Xactimate rate", "Est. hours", "Your rate",
               "Budget wages", "Budget burden", "Budget loaded", "Actual hours", "Actual loaded $", "Variance $"]
    for i, h in enumerate(headers, 1):
        lab.cell(4, i, h)
    _header(lab, 4, 1, 12, fill_navy)
    rows = parsed.labor_rows or []
    for i, row in enumerate(rows):
        r = 5 + i
        lab.cell(r, 1, row.code)
        lab.cell(r, 2, row.description)
        lab.cell(r, 3, row.role).fill = fill_yellow
        lab.cell(r, 3).font = Font(name="Calibri", size=10, color="0000FF")
        lab.cell(r, 4, row.xact_rate).number_format = cur
        lab.cell(r, 5, row.hours).number_format = hrs
        lab.cell(r, 5).fill = fill_yellow
        lab.cell(r, 5).font = Font(name="Calibri", size=10, color="0000FF")
        lab.cell(r, 6, f'=IF(C{r}="Supervisor",Rate_Sup,Rate_Wkr)').number_format = cur
        lab.cell(r, 7, f"=E{r}*F{r}").number_format = cur
        lab.cell(r, 8, f"=G{r}*Burden_Pct").number_format = cur
        lab.cell(r, 9, f"=G{r}+H{r}").number_format = cur
        lab.cell(r, 10).fill = fill_yellow
        lab.cell(r, 10).number_format = hrs
        lab.cell(r, 11).fill = fill_yellow
        lab.cell(r, 11).number_format = cur
        lab.cell(r, 12, f'=IF(OR(J{r}="",K{r}=""),"",K{r}-I{r})').number_format = cur
        for c in range(1, 13):
            lab.cell(r, c).border = thin
            if c not in (3, 5):
                lab.cell(r, c).font = font_body
    last = 4 + max(len(rows), 1)
    if rows:
        last = 4 + len(rows)
        tr = last + 1
    else:
        lab.cell(5, 2, "No labor rows parsed")
        tr = 6
    lab.cell(tr, 1, "TOTAL")
    lab.cell(tr, 5, f"=SUM(E5:E{tr-1})").number_format = hrs
    lab.cell(tr, 7, f"=SUM(G5:G{tr-1})").number_format = cur
    lab.cell(tr, 8, f"=SUM(H5:H{tr-1})").number_format = cur
    lab.cell(tr, 9, f"=SUM(I5:I{tr-1})").number_format = cur
    lab.cell(tr, 10, f"=SUM(J5:J{tr-1})").number_format = hrs
    lab.cell(tr, 11, f"=SUM(K5:K{tr-1})").number_format = cur
    lab.cell(tr, 12, f'=IF(K{tr}=0,"",K{tr}-I{tr})').number_format = cur
    for c in range(1, 13):
        lab.cell(tr, c).fill = fill_navy
        lab.cell(tr, c).font = font_white
        lab.cell(tr, c).border = thin
    dv = DataValidation(type="list", formula1='"Supervisor,Worker"', allow_blank=False)
    lab.add_data_validation(dv)
    if rows:
        dv.add(f"C5:C{last}")
    lab.conditional_formatting.add(f"L5:L{tr}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill_red))
    lab.conditional_formatting.add(f"L5:L{tr}", CellIsRule(operator="lessThan", formula=["0"], fill=fill_green))
    widths = [12, 44, 14, 14, 12, 12, 14, 14, 14, 13, 15, 13]
    for i, w in enumerate(widths, 1):
        lab.column_dimensions[get_column_letter(i)].width = w
    lab.freeze_panes = "A5"
    labor_total_row = tr

    # ----- Materials / Equipment -----
    mat = wb.create_sheet("Materials")
    mat.merge_cells("A1:F1")
    mat["A1"] = "MATERIALS  •  Xactimate material total (edit qty / add rows)"
    mat["A1"].font = font_title
    mat["A1"].fill = fill_navy
    for i, h in enumerate(["Item", "Unit", "Qty", "Unit $", "Budget $", "Actual $", "Variance $"], 1):
        mat.cell(3, i, h)
    _header(mat, 3, 1, 7, fill_teal)
    mat["A4"] = "Xactimate materials (from Internal TAM)"
    mat["B4"] = "LS"
    mat["C4"] = 1
    mat["D4"] = parsed.materials or 0
    mat["D4"].number_format = cur
    mat["D4"].fill = fill_yellow
    mat["E4"] = "=C4*D4"
    mat["E4"].number_format = cur
    mat["F4"].fill = fill_yellow
    mat["F4"].number_format = cur
    mat["G4"] = '=IF(F4="","",F4-E4)'
    mat["G4"].number_format = cur
    for c in range(1, 8):
        mat.cell(4, c).border = thin
    for r in range(5, 12):
        mat.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),0,C{r}*D{r})').number_format = cur
        mat.cell(r, 3).fill = fill_yellow
        mat.cell(r, 4).fill = fill_yellow
        mat.cell(r, 4).number_format = cur
        mat.cell(r, 6).fill = fill_yellow
        mat.cell(r, 6).number_format = cur
        mat.cell(r, 7, f'=IF(F{r}="","",F{r}-E{r})').number_format = cur
        for c in range(1, 8):
            mat.cell(r, c).border = thin
    mat["A12"] = "MATERIALS SUBTOTAL"
    mat["E12"] = "=SUM(E4:E11)"
    mat["F12"] = "=SUM(F4:F11)"
    mat["G12"] = '=IF(F12=0,"",F12-E12)'
    for c in range(1, 8):
        mat.cell(12, c).fill = fill_teal
        mat.cell(12, c).font = font_white
        mat.cell(12, c).border = thin
        if c in (5, 6, 7):
            mat.cell(12, c).number_format = cur
    for i, w in enumerate([48, 10, 10, 12, 14, 14, 14], 1):
        mat.column_dimensions[get_column_letter(i)].width = w

    eq = wb.create_sheet("Equipment & Rentals")
    eq.merge_cells("A1:F1")
    eq["A1"] = "EQUIPMENT & RENTALS  •  Xactimate equipment total"
    eq["A1"].font = font_title
    eq["A1"].fill = fill_navy
    for i, h in enumerate(["Item", "Unit", "Qty", "Unit $", "Budget $", "Actual $", "Variance $"], 1):
        eq.cell(3, i, h)
    _header(eq, 3, 1, 7, PatternFill("solid", fgColor="2E5A88"))
    eq["A4"] = "Xactimate equipment / rentals (from Internal TAM)"
    eq["B4"] = "LS"
    eq["C4"] = 1
    eq["D4"] = parsed.equipment or 0
    eq["D4"].number_format = cur
    eq["D4"].fill = fill_yellow
    eq["E4"] = "=C4*D4"
    eq["E4"].number_format = cur
    eq["F4"].fill = fill_yellow
    eq["F4"].number_format = cur
    eq["G4"] = '=IF(F4="","",F4-E4)'
    eq["G4"].number_format = cur
    for c in range(1, 8):
        eq.cell(4, c).border = thin
    for r in range(5, 12):
        eq.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),0,C{r}*D{r})').number_format = cur
        eq.cell(r, 3).fill = fill_yellow
        eq.cell(r, 4).fill = fill_yellow
        eq.cell(r, 4).number_format = cur
        eq.cell(r, 6).fill = fill_yellow
        eq.cell(r, 6).number_format = cur
        eq.cell(r, 7, f'=IF(F{r}="","",F{r}-E{r})').number_format = cur
        for c in range(1, 8):
            eq.cell(r, c).border = thin
    eq["A12"] = "EQUIPMENT SUBTOTAL"
    eq["E12"] = "=SUM(E4:E11)"
    eq["F12"] = "=SUM(F4:F11)"
    eq["G12"] = '=IF(F12=0,"",F12-E12)'
    for c in range(1, 8):
        eq.cell(12, c).fill = PatternFill("solid", fgColor="2E5A88")
        eq.cell(12, c).font = font_white
        eq.cell(12, c).border = thin
        if c in (5, 6, 7):
            eq.cell(12, c).number_format = cur
    for i, w in enumerate([52, 10, 10, 12, 14, 14, 14], 1):
        eq.column_dimensions[get_column_letter(i)].width = w

    fees = wb.create_sheet("Fees & Other")
    fees["A1"] = "FEES, PERMITS & ROYALTY"
    fees["A1"].font = font_title
    fees["A1"].fill = fill_navy
    fees.merge_cells("A1:E1")
    for i, h in enumerate(["Item", "Budget $", "Actual $", "Variance $", "Notes"], 1):
        fees.cell(3, i, h)
    _header(fees, 3, 1, 5, fill_navy)
    fees["A4"] = "Mobilization / misc. fees (edit)"
    fees["B4"] = 0
    fees["B4"].fill = fill_yellow
    fees["B4"].number_format = cur
    fees["C4"].fill = fill_yellow
    fees["C4"].number_format = cur
    fees["D4"] = '=IF(C4="","",C4-B4)'
    fees["D4"].number_format = cur
    fees["E4"] = "Add dump fees, permits, mobilization if not in equipment"
    fees["A5"] = "Royalty (negotiated price × royalty %)"
    fees["B5"] = "=Billable_Amt*Royalty_Pct"
    fees["B5"].number_format = cur
    fees["C5"].fill = fill_yellow
    fees["C5"].number_format = cur
    fees["D5"] = '=IF(C5="","",C5-B5)'
    fees["D5"].number_format = cur
    fees["E5"] = "Follows Summary!B4"
    for r in (4, 5):
        for c in range(1, 6):
            fees.cell(r, c).border = thin
            fees.cell(r, c).font = font_small if c == 5 else font_body
    fees["A6"] = "SUBTOTAL"
    fees["B6"] = "=B4+B5"
    fees["C6"] = "=SUM(C4:C5)"
    fees["D6"] = '=IF(C6=0,"",C6-B6)'
    for c in range(1, 6):
        fees.cell(6, c).fill = fill_navy
        fees.cell(6, c).font = font_white
        fees.cell(6, c).border = thin
        if c in (2, 3, 4):
            fees.cell(6, c).number_format = cur
    fees.column_dimensions["A"].width = 44
    fees.column_dimensions["E"].width = 50

    act = wb.create_sheet("Actuals Tracker")
    act.merge_cells("A1:F1")
    act["A1"] = "ACTUALS TRACKER  •  invoices & payroll"
    act["A1"].font = font_title
    act["A1"].fill = fill_navy
    for i, h in enumerate(["Date", "Vendor / employee", "Cost type", "Description", "Hours", "Amount $"], 1):
        act.cell(3, i, h)
    _header(act, 3, 1, 6, fill_teal)
    for r in range(4, 34):
        for c in range(1, 7):
            act.cell(r, c).border = thin
            act.cell(r, c).fill = fill_yellow
        act.cell(r, 1).number_format = "YYYY-MM-DD"
        act.cell(r, 5).number_format = hrs
        act.cell(r, 6).number_format = cur
    dv2 = DataValidation(
        type="list",
        formula1='"Labor-Supervisor,Labor-Worker,Materials,Equipment / Rentals,Fees / Permits,Royalty,Tax,Other"',
        allow_blank=True,
    )
    act.add_data_validation(dv2)
    dv2.add("C4:C33")
    act["A34"] = "TOTAL"
    act["E34"] = "=SUM(E4:E33)"
    act["F34"] = "=SUM(F4:F33)"
    act["E34"].number_format = hrs
    act["F34"].number_format = cur
    for c in range(1, 7):
        act.cell(34, c).fill = fill_teal
        act.cell(34, c).font = font_white
        act.cell(34, c).border = thin
    for i, w in enumerate([14, 26, 22, 40, 12, 14], 1):
        act.column_dimensions[get_column_letter(i)].width = w

    # ----- Summary -----
    sm.merge_cells("A1:G1")
    sm["A1"] = f"RESTOPROS RECON TRACKER  •  {job}"
    sm["A1"].font = font_title
    sm["A1"].fill = fill_navy
    sm.row_dimensions[1].height = 28
    sm.merge_cells("A2:G2")
    sm["A2"] = f"Estimate {est}   •   Price list {parsed.price_list or '—'}   •   Built {date.today().isoformat()}   •   Internal TAM recon"
    sm["A2"].font = font_small
    sm["A2"].fill = fill_lt_navy

    sm.merge_cells("A3:G3")
    sm["A3"] = "FINAL NEGOTIATED PRICE  —  royalty and profitability follow this box"
    sm["A3"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    sm["A3"].fill = fill_teal
    sm["A4"] = "Negotiated / collected price"
    sm["A4"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    sm["B4"] = negotiated
    sm["B4"].number_format = cur
    sm["B4"].font = font_input
    sm["B4"].fill = fill_yellow
    sm["B4"].alignment = Alignment(horizontal="center", vertical="center")
    sm["C4"] = "Xactimate RCV (reference)"
    sm["D4"] = rcv
    sm["D4"].number_format = cur
    sm["E4"] = "vs RCV"
    sm["F4"] = '=IF(D4=0,"",B4-D4)'
    sm["F4"].number_format = cur
    sm["G4"] = '=IF(D4=0,"",(B4-D4)/D4)'
    sm["G4"].number_format = pct
    for c in range(1, 8):
        sm.cell(4, c).border = thin
        sm.cell(4, c).fill = fill_lt_gold
    sm["B4"].fill = fill_yellow
    sm.row_dimensions[4].height = 28
    sm.merge_cells("A5:G5")
    sm["A5"] = "Yellow B4 is the single driver for royalty and revenue. Change rates on Assumptions."
    sm["A5"].font = font_small

    sm["A7"] = "COST COMPARISON"
    sm["A7"].font = font_h2
    for i, h in enumerate(["Cost bucket", "Xactimate $", "Internal budget $", "Actual $", "Var $", "Var %", "Notes"], 1):
        sm.cell(8, i, h)
    _header(sm, 8, 1, 7, fill_navy)

    sm["A9"] = "Labor LOADED"
    sm["B9"] = parsed.labor_loaded or parsed.labor_direct
    sm["C9"] = f"='Labor Budget'!I{labor_total_row}"
    sm["D9"] = f"=IF('Labor Budget'!K{labor_total_row}>0,'Labor Budget'!K{labor_total_row},0)"
    sm["A10"] = "Materials"
    sm["B10"] = parsed.materials
    sm["C10"] = "=Materials!E12"
    sm["D10"] = "=IF(Materials!F12>0,Materials!F12,0)"
    sm["A11"] = "Equipment & rentals"
    sm["B11"] = parsed.equipment
    sm["C11"] = "='Equipment & Rentals'!E12"
    sm["D11"] = "=IF('Equipment & Rentals'!F12>0,'Equipment & Rentals'!F12,0)"
    sm["A12"] = "Other fees"
    sm["B12"] = 0
    sm["C12"] = "='Fees & Other'!B4"
    sm["D12"] = "=IF('Fees & Other'!C4<>\"\",'Fees & Other'!C4,0)"
    sm["A13"] = "Royalty (negotiated × rate)"
    sm["B13"] = 0
    sm["C13"] = "=Billable_Amt*Royalty_Pct"
    sm["D13"] = "=IF('Fees & Other'!C5<>\"\",'Fees & Other'!C5,0)"
    sm["A14"] = "Contingency"
    sm["B14"] = 0
    sm["C14"] = "=(C10+C11)*Contingency_Pct"
    sm["D14"] = 0
    sm["D14"].fill = fill_yellow
    for r in range(9, 15):
        sm.cell(r, 2).number_format = cur
        sm.cell(r, 3).number_format = cur
        sm.cell(r, 4).number_format = cur
        sm.cell(r, 5, f'=IF(D{r}=0,"",D{r}-C{r})').number_format = cur
        sm.cell(r, 6, f'=IF(OR(C{r}=0,E{r}=""),"",E{r}/C{r})').number_format = pct
        for c in range(1, 8):
            sm.cell(r, c).border = thin
            sm.cell(r, c).font = font_body
        sm.row_dimensions[r].height = 22
    sm["A15"] = "JOB COST TOTAL"
    sm["B15"] = "=B9+B10+B11+B12+B13+B14"
    sm["C15"] = "=C9+C10+C11+C12+C13+C14"
    sm["D15"] = "=D9+D10+D11+D12+D13+D14"
    sm["E15"] = '=IF(D15=0,"",D15-C15)'
    sm["F15"] = '=IF(OR(C15=0,E15=""),"",E15/C15)'
    for c in range(1, 8):
        sm.cell(15, c).fill = fill_lt_navy
        sm.cell(15, c).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        sm.cell(15, c).border = thin
        if c in (2, 3, 4, 5):
            sm.cell(15, c).number_format = cur
        if c == 6:
            sm.cell(15, c).number_format = pct

    sm.conditional_formatting.add("E9:E15", CellIsRule(operator="greaterThan", formula=["0"], fill=fill_red))
    sm.conditional_formatting.add("E9:E15", CellIsRule(operator="lessThan", formula=["0"], fill=fill_green))

    sm.merge_cells("A17:G17")
    sm["A17"] = "PROFITABILITY  —  revenue = negotiated price B4"
    sm["A17"].font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    sm["A17"].fill = fill_navy
    for i, h in enumerate(["Metric", "Budget", "Actual (or budget if blank)", "Notes"], 1):
        sm.cell(18, i, h)
    _header(sm, 18, 1, 4, fill_teal)
    sm["A19"] = "Billable revenue"
    sm["B19"] = "=Billable_Amt"
    sm["C19"] = "=Billable_Amt"
    sm["D19"] = "Summary!B4"
    sm["A20"] = "Job cost"
    sm["B20"] = "=C15"
    sm["C20"] = "=IF(D15=0,C15,D15)"
    sm["A21"] = "Royalty"
    sm["B21"] = "=C13"
    sm["C21"] = "=IF(D13=0,C13,D13)"
    sm["A22"] = "GROSS PROFIT $"
    sm["B22"] = "=B19-B20"
    sm["C22"] = "=C19-C20"
    sm["A23"] = "Gross margin %"
    sm["B23"] = '=IF(B19=0,"",B22/B19)'
    sm["C23"] = '=IF(C19=0,"",C22/C19)'
    sm["A24"] = "Profit if contingency not spent"
    sm["B24"] = "=B19-(C9+C10+C11+C12+C13)"
    sm["C24"] = "=C19-(IF(D9=0,C9,D9)+IF(D10=0,C10,D10)+IF(D11=0,C11,D11)+IF(D12=0,C12,D12)+IF(D13=0,C13,D13))"
    for r in range(19, 25):
        for c in range(1, 5):
            sm.cell(r, c).border = thin
        for c in (2, 3):
            sm.cell(r, c).number_format = pct if r == 23 else cur
        sm.row_dimensions[r].height = 22
    for c in range(1, 5):
        sm.cell(22, c).fill = fill_lt_gold
        sm.cell(22, c).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        sm.cell(23, c).fill = fill_lt_gold
        sm.cell(23, c).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    sm.conditional_formatting.add("B22:C22", CellIsRule(operator="lessThan", formula=["0"], fill=fill_red))
    sm.conditional_formatting.add("B22:C22", CellIsRule(operator="greaterThan", formula=["0"], fill=fill_green))

    sm.merge_cells("A26:B26")
    sm["A26"] = "BUDGET GROSS PROFIT"
    sm["A26"].font = font_white
    sm["A26"].fill = fill_navy
    sm.merge_cells("C26:D26")
    sm["C26"] = "=B22"
    sm["C26"].number_format = cur
    sm["C26"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    sm["C26"].fill = fill_lt_gold
    sm["E26"] = "MARGIN"
    sm["E26"].font = font_white
    sm["E26"].fill = fill_teal
    sm.merge_cells("F26:G26")
    sm["F26"] = "=B23"
    sm["F26"].number_format = pct
    sm["F26"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    sm["F26"].fill = fill_lt_teal
    sm.row_dimensions[26].height = 26

    sm.merge_cells("A28:G29")
    sm["A28"] = (
        "Source file must be an Xactimate PDF Internal TAM report with all print selections checked: "
        "Coversheet, Line item detail, Summary, Recap by room, Recap by category, Labor breakdown, Sketch. "
        "CLN-S hours priced at supervisor rate; all other codes at worker rate. Yellow cells are inputs."
    )
    sm["A28"].font = font_small
    sm["A28"].alignment = Alignment(wrap_text=True, vertical="top")

    for i, w in enumerate([40, 18, 26, 16, 14, 12, 36], 1):
        sm.column_dimensions[get_column_letter(i)].width = w
    sm.freeze_panes = "A8"
    sm.page_setup.orientation = "landscape"
    sm.page_setup.paperSize = sm.PAPERSIZE_TABLOID
    sm.page_setup.fitToWidth = 1
    sm.page_setup.fitToHeight = 1
    sm.sheet_properties.pageSetUpPr.fitToPage = True
    sm.oddHeader.left.text = "Restopros Recon Tracker"
    sm.oddFooter.right.text = "Page &P of &N"

    # sheet order
    order = ["Summary", "Assumptions", "Labor Budget", "Materials", "Equipment & Rentals", "Fees & Other", "Actuals Tracker"]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    wb.save(path)
    return path
